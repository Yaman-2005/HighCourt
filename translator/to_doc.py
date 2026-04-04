import os
from fpdf import FPDF
from fpdf.enums import XPos, YPos
from docx import Document
from docx.shared import Pt

def create_pdf(input_txt, output_pdf, font_path):
    print(f"[*] Creating PDF: {output_pdf}")
    pdf = FPDF()
    
    # Enable complex text shaping (Mandatory for Bengali)
    try:
        pdf.set_text_shaping(True)
    except:
        print("[!] uharfbuzz not found. Bengali will look like gibberish.")

    pdf.add_page()
    
    # Check for font
    if not os.path.exists(font_path):
        print(f"[!] Font not found at {font_path}. Please check the path.")
        return

    pdf.add_font("BengaliFont", style="", fname=font_path)
    pdf.set_font("BengaliFont", size=11)

    with open(input_txt, "r", encoding="utf-8") as f:
        for line in f:
            content = line.strip()
            if not content:
                pdf.ln(5)
                continue

            # 1. Handle Margin Markers [ক], [খ]...
            if content.startswith("[") and content.endswith("]"):
                pdf.set_font("BengaliFont", size=12)
                pdf.set_text_color(120, 120, 120) # Gray
                pdf.cell(0, 10, text=content, new_x=XPos.LMARGIN, new_y=YPos.NEXT, align='L')
                pdf.set_text_color(0, 0, 0) # Back to black
            
            # 2. Handle Headers
            elif any(x in content.upper() for x in ["SUPREME COURT", "REPORTS", "PAGE"]):
                pdf.set_font("BengaliFont", size=10)
                pdf.cell(0, 10, text=content, new_x=XPos.LMARGIN, new_y=YPos.NEXT, align='C')
            
            # 3. Handle Body Text
            else:
                pdf.set_font("BengaliFont", size=11)
                pdf.multi_cell(0, 8, text=content)
                pdf.ln(2)

    pdf.output(output_pdf)
    print(f"[+] PDF Generated: {os.path.abspath(output_pdf)}")

def create_docx(input_txt, output_docx):
    print(f"[*] Creating DOCX: {output_docx}")
    doc = Document()
    
    with open(input_txt, "r", encoding="utf-8") as f:
        for line in f:
            content = line.strip()
            if not content:
                continue
            
            p = doc.add_paragraph(content)
            # You can set the font name here (e.g., 'Nikosh' or 'Vrinda')
            # but Word usually handles Bengali Unicode automatically.
            run = p.runs[0]
            run.font.size = Pt(11)

    doc.save(output_docx)
    print(f"[+] DOCX Generated: {os.path.abspath(output_docx)}")

if __name__ == "__main__":
    # CONFIGURATION
    SOURCE_TEXT = "structured_bengali_judgment.txt"
    FONT_FILE = "Kalpurush.ttf" # Ensure this is in your folder
    
    if os.path.exists(SOURCE_TEXT):
        create_pdf(SOURCE_TEXT, "Quick_Test_Output.pdf", FONT_FILE)
        create_docx(SOURCE_TEXT, "Quick_Test_Output.docx")
    else:
        print(f"[!] Error: {SOURCE_TEXT} not found. Check your file name.")